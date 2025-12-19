"""
Quick test script to verify python-docx installation and docx_parser functionality
"""
import sys
import os

print("Testing python-docx installation...")
print("-" * 50)

# Test 1: Import check
try:
    from docx import Document
    print("✓ python-docx module imported successfully")
except ImportError as e:
    print(f"✗ Failed to import python-docx: {e}")
    sys.exit(1)

# Test 2: File exists check
docx_file = 'inputs/docx/multimodal_llm_overview.docx'
if os.path.exists(docx_file):
    print(f"✓ Input file exists: {docx_file}")
else:
    print(f"✗ Input file not found: {docx_file}")
    sys.exit(1)

# Test 3: Load document
try:
    doc = Document(docx_file)
    print(f"✓ Document loaded successfully")
    print(f"  - Paragraphs: {len(doc.paragraphs)}")
    print(f"  - First paragraph: {doc.paragraphs[0].text[:50] if doc.paragraphs else 'N/A'}...")
except Exception as e:
    print(f"✗ Failed to load document: {e}")
    sys.exit(1)

# Test 4: Import docx_parser
try:
    sys.path.insert(0, 'src')
    from docx_parser import generate_markdown_from_docx
    print("✓ docx_parser module imported successfully")
except ImportError as e:
    print(f"✗ Failed to import docx_parser: {e}")
    sys.exit(1)

print("-" * 50)
print("✅ All tests passed! python-docx is ready to use.")
print("\nYou can now run:")
print("  python src/docx_parser.py")

